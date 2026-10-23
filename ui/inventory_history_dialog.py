
from PySide6.QtWidgets import (
    QDialog,
    QVBoxLayout,
    QLabel,
    QTableWidget,
    QPushButton,
    QTableWidgetItem,
    QLineEdit,
    QComboBox,
    QDateEdit,
    QMessageBox,
    QHBoxLayout,
    QFileDialog
)
from PySide6.QtCore import QDate
from database.movimiento_repository import obtener_movimientos
from datetime import datetime
import csv
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment
from openpyxl.utils import get_column_letter

class InventoryHistoryDialog(QDialog):
    def __init__(self, parent=None):
        super().__init__(parent)

        self.setWindowTitle("Historial de movimientos de inventario")
        self.resize(1000, 500)

        layout = QVBoxLayout(self)

        titulo = QLabel("Historial de movimientos de inventario")
        layout.addWidget(titulo)

        filtros_layout = QHBoxLayout()
        self.busqueda = QLineEdit()
        self.busqueda.setPlaceholderText("Buscar por codigo o por nombre...")
        filtros_layout.addWidget(self.busqueda)
        titulo_tipo = QLabel("Tipo:")
        filtros_layout.addWidget(titulo_tipo)
        self.filtro_tipo = QComboBox()
        self.filtro_tipo.addItems([
            "TODOS",
            "ENTRADA",
            "SALIDA",
            "AJUSTE",
        ])
        self.filtro_tipo.setCurrentIndex(0)
        filtros_layout.addWidget(self.filtro_tipo)

        self.fecha_desde = QDateEdit()
        self.fecha_desde.setCalendarPopup(True)
        self.fecha_hasta = QDateEdit()
        self.fecha_hasta.setCalendarPopup(True)
        self.fecha_desde.setDate(QDate(2000,1,1))
        self.fecha_hasta.setDate(QDate.currentDate())

        filtros_layout.addWidget(QLabel("Desde:"))
        filtros_layout.addWidget(self.fecha_desde)
        filtros_layout.addWidget(QLabel("Hasta:"))
        filtros_layout.addWidget(self.fecha_hasta)

        self.btn_limpiar_filtro = QPushButton("Limpiar")
        self.btn_limpiar_filtro.clicked.connect(self.limpiar_filtro)
        filtros_layout.addWidget(self.btn_limpiar_filtro)

        self.btn_exportar = QPushButton("Eportar CSV")
        self.btn_exportar.clicked.connect(self.exportar_csv)
        filtros_layout.addWidget(self.btn_exportar)

        self.btn_exportar_excel = QPushButton("Exportar Excel")
        self.btn_exportar_excel.clicked.connect(self.exportar_excel)
        filtros_layout.addWidget(self.btn_exportar_excel)

        layout.addLayout(filtros_layout)

        self.tabla_movimientos = QTableWidget()
        self.tabla_movimientos.setColumnCount(8)
        self.tabla_movimientos.setHorizontalHeaderLabels([
            "Fecha",
            "Código",
            "Producto",
            "Tipo",
            "Cantidad",
            "Stock anterior",
            "Stock nuevo",
            "Observación",
        ])

        layout.addWidget(self.tabla_movimientos)

        self.btn_cerrar = QPushButton("Cerrar")
        self.btn_cerrar.clicked.connect(self.accept)
        layout.addWidget(self.btn_cerrar)

        self.cargar_movimientos()

        self.busqueda.textChanged.connect(self.filtrar_inventario)
        self.filtro_tipo.currentTextChanged.connect(self.filtrar_inventario)
        self.fecha_desde.dateChanged.connect(self.filtrar_inventario)
        self.fecha_hasta.dateChanged.connect(self.filtrar_inventario)

    def cargar_movimientos(self):
        movimientos = obtener_movimientos()

        self.tabla_movimientos.setRowCount(0)

        for movimiento in movimientos:
            fila = self.tabla_movimientos.rowCount()
            self.tabla_movimientos.insertRow(fila)

            producto = movimiento.producto

            valores = [
                movimiento.fecha.strftime("%d/%m/%Y %H:%M"),
                producto.codigo,
                producto.nombre,
                movimiento.tipo,
                str(movimiento.cantidad),
                (
                    str(movimiento.stock_anterior)
                    if movimiento.stock_anterior is not None
                    else "-"
                ),
                (
                    str(movimiento.stock_nuevo)
                    if movimiento.stock_nuevo is not None
                    else "-"
                ), 
                movimiento.observacion or "-"
            ]

            for columna, valor in enumerate(valores):
                self.tabla_movimientos.setItem(
                    fila, columna, QTableWidgetItem(valor)
                )

    def filtrar_inventario(self, texto=None):
        emisor = self.sender()
        texto = self.busqueda.text().strip().lower()
        tipo_seleccionado = self.filtro_tipo.currentText()

        fecha_desde = self.fecha_desde.date()
        fecha_hasta = self.fecha_hasta.date()
        rango_fechas_valido = fecha_hasta >= fecha_desde

        if not rango_fechas_valido and emisor in(self.fecha_desde, self.fecha_hasta):
            QMessageBox.warning(
                self,
                "Rango de fechas inválido",
                "La fecha hasta debe ser mayor o igual a la fecha desde."
            )

        fecha_desde = datetime(
            fecha_desde.year(),
            fecha_desde.month(),
            fecha_desde.day(),
            0,0,0
        )

        fecha_hasta = datetime(
            fecha_hasta.year(),
            fecha_hasta.month(),
            fecha_hasta.day(),
            23,59,59
        )
        
        for fila in range(self.tabla_movimientos.rowCount()):
            codigo = self.tabla_movimientos.item(fila,1).text().lower()
            nombre = self.tabla_movimientos.item(fila,2).text().lower()
            tipo = self.tabla_movimientos.item(fila,3).text()
            fecha_movimiento = datetime.strptime(
                self.tabla_movimientos.item(fila,0).text(),
                "%d/%m/%Y %H:%M"
            )

            coincide_producto = (
                texto in codigo or texto in nombre
            )

            coincide_tipo = (
                tipo_seleccionado == "TODOS" or tipo == tipo_seleccionado
            )

            coincide_fecha =( not rango_fechas_valido 
                            or  fecha_desde <= fecha_movimiento <= fecha_hasta
            )
            mostrar_fila = coincide_producto and coincide_tipo and coincide_fecha
            
            self.tabla_movimientos.setRowHidden(fila, not mostrar_fila)    

    def limpiar_filtro(self):
        self.filtro_tipo.setCurrentIndex(0)
        self.busqueda.clear()
        self.fecha_desde.setDate(QDate(2000,1,1))
        self.fecha_hasta.setDate(QDate.currentDate())

    def exportar_csv(self):
        ruta, _=QFileDialog.getSaveFileName(
            self,
            "Guardar historial de inventario",
            "historial_inventario.csv",
            "Archivos CSV (*.csv)"
        )

        if not ruta:
            return

        if not ruta.lower().endswith(".csv"):
            ruta += ".csv"

        with open(ruta, "w", newline = "", encoding="utf-8-sig") as archivo:
            escritor = csv.writer(archivo, delimiter=";")

            encabezados = [
                self.tabla_movimientos.horizontalHeaderItem(col).text()
                for col in range(self.tabla_movimientos.columnCount())
            ]

            escritor.writerow(encabezados)

            for fila in range(self.tabla_movimientos.rowCount()):
                if self.tabla_movimientos.isRowHidden(fila):
                    continue

                valores = [
                    self.tabla_movimientos.item(fila, col).text()
                    if self.tabla_movimientos.item(fila, col) is not None
                    else ""
                    for col in range(self.tabla_movimientos.columnCount())
                ]

                escritor.writerow(valores)

        QMessageBox.information(
            self,
            "Exportación completada",
            f"El historial se guardó en:\n{ruta}"
        )

    def exportar_excel(self):
        ruta, _=QFileDialog.getSaveFileName(
            self,
            "Guardar historial de inventario",
            "historial_inventario.xlsx",
            "Archivos Excel (*.xlsx)"
        )

        if not ruta:
            return

        if not ruta.lower().endswith(".xlsx"):
            ruta += ".xlsx"

        libro = Workbook()
        hoja = libro.active
        hoja.title = "Historial de inventario"

        encabezados = [
            self.tabla_movimientos.horizontalHeaderItem(col).text()
            for col in range(self.tabla_movimientos.columnCount())
        ]

        hoja.append(encabezados)

        for fila in range(self.tabla_movimientos.rowCount()):
            if self.tabla_movimientos.isRowHidden(fila):
                continue

            valores = [
                self.tabla_movimientos.item(fila, col).text()
                if self.tabla_movimientos.item(fila, col) is not None
                else ""
                for col in range(self.tabla_movimientos.columnCount())
            ]

            hoja.append(valores)

        for celda in hoja[1]:
            celda.font = Font(bold=True, color="FFFFFF")
            celda.fill = PatternFill(
                fill_type="solid",
                fgColor="1F4E78"
            )
            celda.alignment = Alignment(horizontal="center")

        for columna in hoja.columns:
            longitud_maxima = max(
                len(str(celda.value)) if celda.value is not None else 0
                for celda in columna
            )

            letra_columna = get_column_letter(columna[0].column)
            hoja.column_dimensions[letra_columna].width = min(
                longitud_maxima + 3,
                40
            )

        libro.save(ruta)

        QMessageBox.information(
            self,
            "Exportación completada",
            f"El historial se guardó en: \n{ruta}"
        )